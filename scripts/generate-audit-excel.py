import json
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Define path
base_dir = r"d:\[01] PROJECT\HERO"
json_path = os.path.join(base_dir, ".tmp_pages.json")
excel_path = os.path.join(base_dir, "audit_dan_rencana_perbaikan.xlsx")

with open(json_path, 'r', encoding='utf-8') as f:
    pages = json.load(f)

# Sort pages by section and title
pages.sort(key=lambda x: (x.get('section', ''), x.get('title', '')))

# We will build three sheets:
# 1. Daftar Halaman & Status (List of Pages and Statuses)
# 2. Temuan Audit (Audit Findings)
# 3. Rencana Perbaikan (Remediation Plan)

# Sheet 1 data
sheet1_data = []
for p in pages:
    url = p.get('url', '')
    title = p.get('title', '')
    section = p.get('section', '')
    
    # Analyze if it needs approvals/reminders
    needs_approval = "No"
    needs_reminder = "No"
    approval_status = "N/A"
    reminder_status = "N/A"
    
    # Map known pages
    url_lower = url.lower()
    if "activity" in url_lower or "daily-activity" in url_lower:
        needs_approval = "Yes"
        needs_reminder = "Yes"
        approval_status = "Terintegrasi Penuh (Form Template-driven)"
        reminder_status = "SLA & Persistence Aktif (Worker Pending)"
    elif "overtime" in url_lower or "lembur" in url_lower:
        needs_approval = "Yes"
        needs_reminder = "Yes"
        approval_status = "Terintegrasi Penuh"
        reminder_status = "SLA & Persistence Aktif (Worker Pending)"
    elif "leave" in url_lower or "izin" in url_lower or "cuti" in url_lower:
        needs_approval = "Yes"
        needs_reminder = "Yes"
        approval_status = "Fondasi Siap (Generic Engine)"
        reminder_status = "Fondasi Siap (SLA Persistence)"
    elif "report" in url_lower or "daily-report" in url_lower:
        needs_approval = "Yes"
        needs_reminder = "Yes"
        approval_status = "Fondasi Siap (Generic Engine)"
        reminder_status = "Fondasi Siap"
    elif "hse" in url_lower or "incident" in url_lower or "hazard" in url_lower:
        needs_approval = "Yes"
        needs_reminder = "Yes"
        approval_status = "Fondasi Siap"
        reminder_status = "Fondasi Siap"
    elif "approval" in url_lower:
        needs_approval = "N/A (Inbox)"
        needs_reminder = "N/A"
        approval_status = "Pusat Approval Workbench Aktif"
        reminder_status = "N/A"
    elif "request-center" in url_lower:
        needs_approval = "N/A (Tracker)"
        needs_reminder = "N/A"
        approval_status = "Request Center Tracking Aktif"
        reminder_status = "N/A"
        
    sheet1_data.append({
        "ID Halaman": p.get('id'),
        "Modul/Section": section,
        "Nama Halaman": title,
        "URL Path": url,
        "Butuh Approval": needs_approval,
        "Butuh Reminder": needs_reminder,
        "Status Integrasi Approval": approval_status,
        "Status Integrasi Reminder": reminder_status
    })

df1 = pd.DataFrame(sheet1_data)

# Sheet 2 data (Temuan Audit)
sheet2_data = [
    {
        "No": 1,
        "Modul / Halaman": "Daily Activity",
        "Temuan Audit": "Integrasi workflow approval dan form template-driven sudah aktif penuh, namun scheduler pengiriman notifikasi reminder belum berjalan otomatis di background.",
        "Dampak": "User / approver tidak menerima email/in-app alert reminder jika tugas approval mendekati batas SLA.",
        "Tingkat Bahaya": "Medium",
        "Status": "Belum Selesai"
    },
    {
        "No": 2,
        "Modul / Halaman": "Approval Inbox",
        "Temuan Audit": "Inbox workbench (/dashboard/approval) sudah menampilkan antrean, durasi SLA, route snapshot, timeline comment, dan tombol aksi cepat. Integrasi data sudah sempurna.",
        "Dampak": "User bisa melakukan Approve/Reject/Revisi secara terpusat.",
        "Tingkat Bahaya": "Low",
        "Status": "Selesai"
    },
    {
        "No": 3,
        "Modul / Halaman": "Overtime Requests",
        "Temuan Audit": "Kalkulasi overtimeMinutes dan penentuan priority point sudah terintegrasi dengan Approval Engine berdasarkan threshold.",
        "Dampak": "Alur pengajuan lembur tervalidasi dengan benar.",
        "Tingkat Bahaya": "Low",
        "Status": "Selesai"
    },
    {
        "No": 4,
        "Modul / Halaman": "Leave / Permission (Cuti/Izin)",
        "Temuan Audit": "Fondasi tabel dan schema approval untuk pengajuan cuti/izin sudah siap, namun form pengajuan belum sepenuhnya dimigrasikan ke template-driven form engine.",
        "Dampak": "Pengajuan cuti belum bisa dinikmati dengan dynamic approval matrix.",
        "Tingkat Bahaya": "Medium",
        "Status": "Dalam Pengembangan"
    },
    {
        "No": 5,
        "Modul / Halaman": "Notification Bell (Header)",
        "Temuan Audit": "Bell notifikasi di header layout dashboard masih menggunakan data mock statis dan belum membaca data riwayat pengiriman dari tabel `notificationEvents`.",
        "Dampak": "User tidak mengetahui secara langsung jika ada tugas baru tanpa membuka halaman inbox approval.",
        "Tingkat Bahaya": "Medium",
        "Status": "Belum Selesai"
    }
]
df2 = pd.DataFrame(sheet2_data)

# Sheet 3 data (Rencana Perbaikan)
sheet3_data = [
    {
        "No": 1,
        "Modul Sasaran": "Notification Bell",
        "Rencana Aksi Perbaikan": "Ganti Bell notifikasi mock dengan component dinamis yang melakukan fetch ke API route `notificationEvents` untuk user aktif.",
        "Prioritas": "High",
        "Target Penyelesaian": "Minggu Depan"
    },
    {
        "No": 2,
        "Modul Sasaran": "Reminder Scheduler",
        "Rencana Aksi Perbaikan": "Buat background cron job worker menggunakan Node scheduler / pg-boss untuk memproses antrean tabel `reminderJobs` secara berkala.",
        "Prioritas": "High",
        "Target Penyelesaian": "2 Minggu"
    },
    {
        "No": 3,
        "Modul Sasaran": "Leave & Permission Form",
        "Rencana Aksi Perbaikan": "Migrasikan UI Form Izin/Cuti ke engine form dynamic agar secara otomatis menggunakan `resolveApprovalRouteForActivity` saat disubmit.",
        "Prioritas": "Medium",
        "Target Penyelesaian": "3 Minggu"
    }
]
df3 = pd.DataFrame(sheet3_data)

# Write Excel using Pandas and Openpyxl for formatting
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    df1.to_excel(writer, sheet_name='Daftar Halaman & Status', index=False)
    df2.to_excel(writer, sheet_name='Temuan Audit', index=False)
    df3.to_excel(writer, sheet_name='Rencana Perbaikan', index=False)

# Stylize Excel
wb = openpyxl.load_workbook(excel_path)
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # Auto-adjust column width
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)
        
    # Style header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    
    # Add border and wrap text to all data cells
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 20
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=11)
            # Alignment rules
            if sheet_name == 'Daftar Halaman & Status':
                if col in [1, 5, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif sheet_name == 'Temuan Audit':
                if col in [1, 5, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif sheet_name == 'Rencana Perbaikan':
                if col in [1, 4]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Save workbook
wb.save(excel_path)
print(f"Generated formatted Excel file at {excel_path}")
