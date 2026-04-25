import csv
import json
from collections import defaultdict
from pathlib import Path

import openpyxl


ROOT = Path(r"D:\[01] PROJECT\HERO")
WORKBOOK = ROOT / "DATA HERO.xlsx"
OUTDIR = ROOT / "tools" / "data_hero_normalized"


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def slug(text):
    return "".join(ch.lower() if ch.isalnum() else "_" for ch in text).strip("_")


def write_csv(path: Path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    raw_header = [clean(v) for v in rows[0]]
    header = []
    seen = defaultdict(int)
    for col in raw_header:
        base = col or "unnamed"
        seen[base] += 1
        header.append(base if seen[base] == 1 else f"{base} #{seen[base]}")

    records = []
    for row in rows[1:]:
        if not any(clean(v) is not None for v in row):
            continue
        rec = {}
        for idx, col in enumerate(header):
            rec[col] = clean(row[idx] if idx < len(row) else None)
        records.append(rec)

    master_department = []
    master_section = []
    master_site = []
    master_work_location = []
    master_job_level = []
    master_gender = []
    master_age_band = []
    master_service_band = []
    master_education = []
    master_employee_status = []
    master_position = []
    org_nodes = []
    users = []

    dept_ids = {}
    section_ids = {}
    site_ids = {}
    work_location_ids = {}
    position_ids = {}

    for dept in sorted({r["Department"] for r in records if r.get("Department")}):
        dept_id = f"DEPT_{slug(dept)}"
        dept_ids[dept] = dept_id
        master_department.append({"department_id": dept_id, "department_name": dept})

    for dept, section in sorted({(r.get("Department"), r.get("Section")) for r in records if r.get("Department") and r.get("Section")}):
        section_id = f"SEC_{slug(dept)}_{slug(section)}"
        section_ids[(dept, section)] = section_id
        master_section.append(
            {
                "section_id": section_id,
                "department_id": dept_ids[dept],
                "section_name": section,
            }
        )

    for site in sorted({r["Location / Site"] for r in records if r.get("Location / Site")}):
        site_id = f"SITE_{slug(site)}"
        site_ids[site] = site_id
        master_site.append({"site_id": site_id, "site_name": site})

    for work in sorted({r["Lokasi Kerja"] for r in records if r.get("Lokasi Kerja")}):
        work_id = f"WORK_{slug(work)}"
        work_location_ids[work] = work_id
        master_work_location.append({"work_location_id": work_id, "work_location_name": work})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Job Level Code"), r.get("Job Level"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_job_level.append({"job_level_code": pair[0], "job_level_name": pair[1]})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Gender Code"), r.get("Gender"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_gender.append({"gender_code": pair[0], "gender_name": pair[1]})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Age Code"), r.get("Age"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_age_band.append({"age_band_code": pair[0], "age_band_name": pair[1]})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Length of Service Code"), r.get("Length of Service"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_service_band.append({"service_band_code": pair[0], "service_band_name": pair[1]})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Last Formal Education Code"), r.get("Last Formal Education"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_education.append({"education_code": pair[0], "education_name": pair[1]})

    seen_pairs = set()
    for r in records:
        pair = (r.get("Demographic Employee Status Code"), r.get("Demographic Employee Status"))
        if pair not in seen_pairs and pair[0] and pair[1]:
            seen_pairs.add(pair)
            master_employee_status.append({"employee_status_code": pair[0], "employee_status_name": pair[1]})

    for level, pangkat, job_level, emp_status in sorted(
        {(r.get("Level"), r.get("Pangkat"), r.get("Job Level"), r.get("Employee Status")) for r in records if r.get("Level") and r.get("Pangkat")},
        key=lambda x: tuple(v or "" for v in x),
    ):
        pos_id = f"POS_{slug(level)}_{slug(pangkat)}_{slug(job_level or 'na')}_{slug(emp_status or 'na')}"
        position_ids[(level, pangkat, job_level, emp_status)] = pos_id
        master_position.append(
            {
                "position_id": pos_id,
                "level_name": level,
                "rank_name": pangkat,
                "job_level_name": job_level,
                "employee_status_name": emp_status,
            }
        )

    org_nodes.append({"org_node_id": "ORG_ROOT", "parent_org_node_id": "", "node_type": "company", "node_name": "HERO", "ref_id": "", "path": "HERO"})
    for dept in master_department:
        dept_node = f"ORG_{dept['department_id']}"
        org_nodes.append({"org_node_id": dept_node, "parent_org_node_id": "ORG_ROOT", "node_type": "department", "node_name": dept["department_name"], "ref_id": dept["department_id"], "path": f"HERO > {dept['department_name']}"})

    section_site_map = defaultdict(set)
    workloc_map = defaultdict(set)
    for r in records:
        if r.get("Department") and r.get("Section") and r.get("Location / Site"):
            section_site_map[(r["Department"], r["Section"])].add(r["Location / Site"])
        if r.get("Department") and r.get("Section") and r.get("Location / Site") and r.get("Lokasi Kerja"):
            workloc_map[(r["Department"], r["Section"], r["Location / Site"])].add(r["Lokasi Kerja"])

    for dept, section in sorted(section_ids):
        section_node = f"ORG_{section_ids[(dept, section)]}"
        org_nodes.append({"org_node_id": section_node, "parent_org_node_id": f"ORG_{dept_ids[dept]}", "node_type": "section", "node_name": section, "ref_id": section_ids[(dept, section)], "path": f"HERO > {dept} > {section}"})
        for site in sorted(section_site_map[(dept, section)]):
            site_node = f"ORG_{section_ids[(dept, section)]}_{site_ids[site]}"
            org_nodes.append({"org_node_id": site_node, "parent_org_node_id": section_node, "node_type": "site", "node_name": site, "ref_id": site_ids[site], "path": f"HERO > {dept} > {section} > {site}"})
            for work in sorted(workloc_map[(dept, section, site)]):
                work_node = f"ORG_{section_ids[(dept, section)]}_{site_ids[site]}_{work_location_ids[work]}"
                org_nodes.append({"org_node_id": work_node, "parent_org_node_id": site_node, "node_type": "work_location", "node_name": work, "ref_id": work_location_ids[work], "path": f"HERO > {dept} > {section} > {site} > {work}"})

    for r in records:
        dept = r.get("Department")
        section = r.get("Section")
        site = r.get("Location / Site")
        work = r.get("Lokasi Kerja")
        org_node_id = ""
        if dept and section and site and work and (dept, section) in section_ids and site in site_ids and work in work_location_ids:
            org_node_id = f"ORG_{section_ids[(dept, section)]}_{site_ids[site]}_{work_location_ids[work]}"
        users.append(
            {
                "user_id": r["Employee ID"],
                "employee_id": r["Employee ID"],
                "full_name": r.get("Employee Name") or "",
                "email": r.get("Email Address") or "",
                "password_email": r.get("Password Email") or "",
                "department_id": dept_ids.get(dept, ""),
                "section_id": section_ids.get((dept, section), ""),
                "site_id": site_ids.get(site, ""),
                "work_location_id": work_location_ids.get(work, ""),
                "position_id": position_ids.get((r.get("Level"), r.get("Pangkat"), r.get("Job Level"), r.get("Employee Status")), ""),
                "org_node_id": org_node_id,
                "employee_status": r.get("Employee Status") or "",
                "join_date": r.get("Join Date") or "",
                "birth_date": r.get("Tanggal Lahir") or "",
                "gender_code": r.get("Gender Code") or "",
                "age_band_code": r.get("Age Code") or "",
                "service_band_code": r.get("Length of Service Code") or "",
                "education_code": r.get("Last Formal Education Code") or "",
                "demographic_employee_status_code": r.get("Demographic Employee Status Code") or "",
                "location_category_code": r.get("Location Code") or "",
            }
        )

    write_csv(OUTDIR / "master_departments.csv", master_department, ["department_id", "department_name"])
    write_csv(OUTDIR / "master_sections.csv", master_section, ["section_id", "department_id", "section_name"])
    write_csv(OUTDIR / "master_sites.csv", master_site, ["site_id", "site_name"])
    write_csv(OUTDIR / "master_work_locations.csv", master_work_location, ["work_location_id", "work_location_name"])
    write_csv(OUTDIR / "master_job_levels.csv", master_job_level, ["job_level_code", "job_level_name"])
    write_csv(OUTDIR / "master_genders.csv", master_gender, ["gender_code", "gender_name"])
    write_csv(OUTDIR / "master_age_bands.csv", master_age_band, ["age_band_code", "age_band_name"])
    write_csv(OUTDIR / "master_service_bands.csv", master_service_band, ["service_band_code", "service_band_name"])
    write_csv(OUTDIR / "master_educations.csv", master_education, ["education_code", "education_name"])
    write_csv(OUTDIR / "master_employee_statuses.csv", master_employee_status, ["employee_status_code", "employee_status_name"])
    write_csv(OUTDIR / "master_positions.csv", master_position, ["position_id", "level_name", "rank_name", "job_level_name", "employee_status_name"])
    write_csv(OUTDIR / "organization_nodes.csv", org_nodes, ["org_node_id", "parent_org_node_id", "node_type", "node_name", "ref_id", "path"])
    write_csv(OUTDIR / "user_management.csv", users, ["user_id", "employee_id", "full_name", "email", "password_email", "department_id", "section_id", "site_id", "work_location_id", "position_id", "org_node_id", "employee_status", "join_date", "birth_date", "gender_code", "age_band_code", "service_band_code", "education_code", "demographic_employee_status_code", "location_category_code"])

    report = {
        "source_records": len(records),
        "master_counts": {
            "departments": len(master_department),
            "sections": len(master_section),
            "sites": len(master_site),
            "work_locations": len(master_work_location),
            "job_levels": len(master_job_level),
            "genders": len(master_gender),
            "age_bands": len(master_age_band),
            "service_bands": len(master_service_band),
            "educations": len(master_education),
            "employee_statuses": len(master_employee_status),
            "positions": len(master_position),
        },
        "organization_count": len(org_nodes),
        "user_count": len(users),
        "users_with_org_node": sum(1 for u in users if u["org_node_id"]),
        "users_without_org_node": sum(1 for u in users if not u["org_node_id"]),
        "users_with_position": sum(1 for u in users if u["position_id"]),
        "users_without_position": sum(1 for u in users if not u["position_id"]),
        "users_with_email": sum(1 for u in users if u["email"]),
        "users_without_email": sum(1 for u in users if not u["email"]),
    }
    (OUTDIR / "integration_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
