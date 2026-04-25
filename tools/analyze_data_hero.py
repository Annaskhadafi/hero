import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


WORKBOOK = Path(r"D:\[01] PROJECT\HERO\DATA HERO.xlsx")


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text != "" else None


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    output = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header = [clean(v) for v in (rows[0] if rows else [])]
        body = rows[1:] if len(rows) > 1 else []
        nonempty = [r for r in body if any(clean(v) is not None for v in r)]
        normalized = [tuple(clean(v) for v in row[: len(header)]) for row in nonempty]
        row_dupes = sum(count - 1 for count in Counter(normalized).values() if count > 1)
        samples = []
        for row in nonempty[:5]:
            record = {}
            for idx, col in enumerate(header):
                if col is None:
                    continue
                record[col] = clean(row[idx] if idx < len(row) else None)
            samples.append(record)
        records = []
        for row in nonempty:
            record = {}
            for idx, col in enumerate(header):
                if col is None:
                    continue
                record[col] = clean(row[idx] if idx < len(row) else None)
            records.append(record)

        col_stats = []
        for col in header:
            if col is None:
                continue
            values = [r.get(col) for r in records]
            filled = [v for v in values if v is not None]
            col_stats.append(
                {
                    "column": col,
                    "filled": len(filled),
                    "null": len(values) - len(filled),
                    "distinct": len(set(filled)),
                }
            )

        unique_sets = {}
        for col in ["Employee ID", "Email Address"]:
            if col in header:
                vals = [r.get(col) for r in records if r.get(col) is not None]
                unique_sets[col] = {
                    "filled": len(vals),
                    "distinct": len(set(vals)),
                    "duplicates": len(vals) - len(set(vals)),
                }

        taxonomies = {}
        for col in [
            "Department",
            "Section",
            "Location / Site",
            "Lokasi Kerja",
            "Level",
            "Pangkat",
            "Job Level",
            "Employee Status",
        ]:
            if col in header:
                vals = [r.get(col) for r in records if r.get(col) is not None]
                taxonomies[col] = sorted(set(vals))

        dept_section = defaultdict(set)
        section_site = defaultdict(set)
        for r in records:
            dept = r.get("Department")
            sec = r.get("Section")
            site = r.get("Location / Site")
            if dept and sec:
                dept_section[dept].add(sec)
            if sec and site:
                section_site[sec].add(site)

        output.append(
            {
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "header": header,
                "nonempty_rows": len(nonempty),
                "duplicate_rows": row_dupes,
                "column_stats": col_stats,
                "candidate_unique_keys": unique_sets,
                "taxonomies": taxonomies,
                "department_to_sections": {k: sorted(v) for k, v in sorted(dept_section.items())},
                "section_to_sites": {k: sorted(v) for k, v in sorted(section_site.items())},
                "samples": samples,
            }
        )
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
