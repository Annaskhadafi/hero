import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


WORKBOOK = Path(r"D:\[01] PROJECT\HERO\DATA HERO.xlsx")


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(v) for v in rows[0]]
    body = [r for r in rows[1:] if any(clean(v) is not None for v in r)]

    idx = {name: pos for pos, name in enumerate(header) if name is not None and name not in locals().get('idx', {})}
    records = []
    for row in body:
        rec = {}
        for col, pos in idx.items():
            rec[col] = clean(row[pos] if pos < len(row) else None)
        records.append(rec)

    depts = sorted({r["Department"] for r in records if r.get("Department")})
    sections = sorted({r["Section"] for r in records if r.get("Section")})
    sites = sorted({r["Location / Site"] for r in records if r.get("Location / Site")})
    work_locations = sorted({r["Lokasi Kerja"] for r in records if r.get("Lokasi Kerja")})
    positions = sorted(
        {(r.get("Level"), r.get("Pangkat"), r.get("Job Level"), r.get("Employee Status")) for r in records},
        key=lambda x: tuple("" if v is None else v for v in x),
    )

    org_nodes = set()
    user_mapping_missing = []
    section_site_counts = Counter()
    dept_section_counts = Counter()
    position_counts = Counter()
    site_workloc_counts = Counter()
    incomplete_core = 0

    for r in records:
        dept = r.get("Department")
        section = r.get("Section")
        site = r.get("Location / Site")
        workloc = r.get("Lokasi Kerja")
        level = r.get("Level")
        pangkat = r.get("Pangkat")
        if not all([dept, section, workloc]):
            incomplete_core += 1
        if dept:
            org_nodes.add(("department", dept))
        if dept and section:
            org_nodes.add(("section", dept, section))
            dept_section_counts[(dept, section)] += 1
        if dept and section and site:
            org_nodes.add(("site", dept, section, site))
            section_site_counts[(section, site)] += 1
        if dept and section and site and workloc:
            org_nodes.add(("work_location", dept, section, site, workloc))
            site_workloc_counts[(site, workloc)] += 1
        if level or pangkat:
            position_counts[(level, pangkat, r.get("Job Level"), r.get("Employee Status"))] += 1
        if not all([r.get("Employee ID"), dept, section, workloc, level]):
            user_mapping_missing.append(r.get("Employee ID"))

    out = {
        "record_count": len(records),
        "department_count": len(depts),
        "section_count": len(sections),
        "site_count": len(sites),
        "work_location_count": len(work_locations),
        "position_signature_count": len(positions),
        "org_node_count": len(org_nodes),
        "dept_section_combo_count": len(dept_section_counts),
        "section_site_combo_count": len(section_site_counts),
        "site_work_location_combo_count": len(site_workloc_counts),
        "incomplete_core_records": incomplete_core,
        "user_mapping_missing_count": len(user_mapping_missing),
        "top_department_sizes": [
            {"department": k, "employees": v}
            for k, v in Counter(r.get("Department") for r in records).most_common()
        ],
        "top_section_sizes": [
            {"section": k, "employees": v}
            for k, v in Counter(r.get("Section") for r in records).most_common()
        ],
        "position_signatures": [
            {
                "level": k[0],
                "pangkat": k[1],
                "job_level": k[2],
                "employee_status": k[3],
                "employees": v,
            }
            for k, v in position_counts.most_common()
        ],
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
